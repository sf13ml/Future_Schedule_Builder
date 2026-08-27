import io
import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from models import LayoutBlueprint
from typing import Any, List, Dict
from pydantic import BaseModel

# --- HEBREW DETECTION LAYER ---

def contains_hebrew(text: str) -> bool:
    """Returns True if the string contains at least one Hebrew character."""
    if not isinstance(text, str):
        return False
    return bool(re.search(r"[\u0590-\u05FF]", text))

def detect_dataset_direction(schedule_data: Any) -> bool:
    """Scans raw schedule data to detect if any string contains Hebrew characters."""
    if isinstance(schedule_data, dict):
        raw_assignments = schedule_data.get("assignments") or schedule_data.get("schedule") or []
    elif isinstance(schedule_data, list):
        raw_assignments = schedule_data
    else:
        raw_assignments = []

    for item in raw_assignments:
        if isinstance(item, dict):
            for val in item.values():
                if contains_hebrew(str(val)):
                    return True
        elif contains_hebrew(str(item)):
            return True
    return False

# --- DYNAMIC TIME SLOT FORMATTER ---

def standardize_shift_hours(raw_slot: str) -> str:
    """Formats raw slot strings (e.g., '06-10', '14-18') into 'HH:MM-HH:MM' without hardcoding defaults."""
    if not raw_slot or str(raw_slot).strip().lower() in ("none", "", "nan", "unassigned", "main", "standard"):
        return ""
    
    clean_str = str(raw_slot).strip().replace("–", "-")
    parts = clean_str.split("-")
    
    if len(parts) == 2:
        start_p, end_p = parts[0].strip(), parts[1].strip()
        if ":" not in start_p and start_p.isdigit():
            start_p = f"{int(start_p):02d}:00"
        if ":" not in end_p and end_p.isdigit():
            end_p = f"{int(end_p):02d}:00"
        return f"{start_p}-{end_p}"
        
    return clean_str

# --- DATA CONTRACT / NORMALIZATION LAYER ---

class CanonicalAssignment(BaseModel):
    shift_id: str
    day: str
    time_block: str
    time_slot: str  
    job_name: str
    worker_uid: str
    worker_name: str

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "CanonicalAssignment":
        """Normalizes heterogeneous input dictionary into a strict schema contract."""
        
        # 1. Resolve Job Display Name
        job_raw = (
            data.get("job_name")
            or data.get("job_group")
            or data.get("role")
            or data.get("job_category")
            or (data.get("job", {}).get("name") if isinstance(data.get("job"), dict) else data.get("job"))
            or "General Shift"
        )

        # 2. Resolve Time Block & Time Slot separately
        time_slot_raw = data.get("time_slot") or data.get("shift_hours") or ""
        time_block_raw = data.get("time_block") or data.get("block") or "General"

        # 3. Resolve Day Name
        day_raw = data.get("day_name") or data.get("day") or ""

        # 4. Resolve Worker Identifiers
        worker_obj = data.get("worker") if isinstance(data.get("worker"), dict) else {}
        worker_uid_raw = (
            data.get("worker_uid")
            or data.get("worker_id")
            or worker_obj.get("uid")
            or worker_obj.get("id")
            or ""
        )
        worker_name_raw = (
            data.get("worker_name")
            or worker_obj.get("name")
            or (str(data.get("worker")) if data.get("worker") and not isinstance(data.get("worker"), dict) else "")
        )

        clean_worker_name = str(worker_name_raw).strip() if worker_name_raw and str(worker_name_raw).lower() != "none" else ""

        # 5. Resolve Shift ID
        shift_id_raw = data.get("id") or data.get("shift_id") or ""

        return cls(
            shift_id=str(shift_id_raw).strip(),
            day=str(day_raw).strip().capitalize(),
            time_block=str(time_block_raw).strip(),
            time_slot=str(time_slot_raw).strip(),
            job_name=str(job_raw).strip(),
            worker_uid=str(worker_uid_raw).strip(),
            worker_name=clean_worker_name
        )


# --- MAIN EXPORT FUNCTION ---

def generate_excel_schedule(schedule_data: Any, blueprint: LayoutBlueprint = LayoutBlueprint(), is_rtl: bool = None) -> io.BytesIO:
    if is_rtl is None:
        is_rtl = detect_dataset_direction(schedule_data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "לוח משמרות" if is_rtl else "Schedule"
    
    ws.views.sheetView[0].rightToLeft = is_rtl

    try:
        ws.sheet_view.showGridLines = blueprint.show_gridlines
    except AttributeError:
        pass

    # Unpack Payload
    if isinstance(schedule_data, dict):
        raw_assignments = schedule_data.get("assignments") or schedule_data.get("schedule") or []
        metrics = schedule_data.get("metrics") or {}
        diagnostics = schedule_data.get("diagnostics") or schedule_data.get("audit_logs") or schedule_data.get("warnings") or []
        status = schedule_data.get("status", "unknown")
    else:
        raw_assignments = schedule_data if isinstance(schedule_data, list) else []
        metrics = {}
        diagnostics = []
        status = "optimal"

    assignments: List[CanonicalAssignment] = [
        CanonicalAssignment.from_dict(a) for a in raw_assignments if isinstance(a, dict)
    ]

    # Tab 1: Schedule Matrix / Tabular
    if getattr(blueprint, "layout_type", "matrix") == "tabular":
        _render_tabular_layout(ws, assignments, blueprint, is_rtl)
    else:
        _render_matrix_layout(ws, assignments, blueprint, is_rtl)

    # Tab 2: Analytics & Metrics
    if metrics or any(a.worker_name for a in assignments):
        _render_analytics_sheet(wb, metrics, is_rtl)

    # Tab 3: System Diagnostics & Debug
    if not diagnostics and status in ("infeasible", "infeasable"):
        diagnostics = [{
            "severity": "CRITICAL",
            "code": "SOLVER_INFEASIBLE",
            "entity": "Global Matrix",
            "reason": "The model could not be resolved under strict constraints.",
            "remediation": "Check for conflicting hard constraints or manpower shortages."
        }]
    elif not diagnostics:
        diagnostics = [{
            "severity": "INFO",
            "code": "STATUS_OK",
            "entity": "Global Schedule",
            "reason": "Schedule compiled successfully with all constraints satisfied.",
            "remediation": "None required."
        }]

    _render_diagnostics_sheet(wb, diagnostics, is_rtl)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# --- RENDER MATRIX LAYOUT ---

def _render_matrix_layout(ws, assignments: List[CanonicalAssignment], blueprint: LayoutBlueprint, is_rtl: bool):
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10)
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    h_align = "right" if is_rtl else "left"
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_side = Alignment(horizontal=h_align, vertical="center")

    x_headers = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]

    # Dynamic Block Discovery (preserves payload order without hardcoded block assumptions)
    block_order = []
    row_tree = {}

    for a in assignments:
        block_label = str(a.time_block).strip() if a.time_block else "General"
        if block_label not in block_order:
            block_order.append(block_label)

        display_slot = standardize_shift_hours(a.time_slot)
        job_name = str(a.job_name).strip()
        shift_id = str(a.shift_id)

        slot_idx = 1
        parts = shift_id.split("_")
        if parts and parts[-1].isdigit():
            slot_idx = int(parts[-1])

        day_key = str(a.day).capitalize()

        current_s_idx = slot_idx
        while True:
            candidate_key = (block_label, job_name, display_slot, current_s_idx)
            if candidate_key not in row_tree or day_key not in row_tree[candidate_key]:
                row_key = candidate_key
                break
            current_s_idx += 1

        if row_key not in row_tree:
            row_tree[row_key] = {}
        
        row_tree[row_key][day_key] = a.worker_name

    block_rank_map = {b_label: idx for idx, b_label in enumerate(block_order)}
    sorted_row_keys = sorted(
        row_tree.keys(),
        key=lambda k: (block_rank_map.get(k[0], 99), k[1], k[2], k[3])
    )

    headers = ["Time Block", "Job Category", "Shift Hours"] + x_headers
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font, cell.fill, cell.alignment = font_header, fill_header, align_center

    start_row = 2
    for r_idx, (t_block, j_name, t_slot, s_idx) in enumerate(sorted_row_keys, start=start_row):
        c1 = ws.cell(row=r_idx, column=1, value=t_block)
        c2 = ws.cell(row=r_idx, column=2, value=j_name)
        c3 = ws.cell(row=r_idx, column=3, value=t_slot)
        
        for c in (c1, c2, c3):
            c.font, c.fill, c.border = font_body, fill_white, border_cell
        c1.alignment, c3.alignment = align_center, align_center
        c2.alignment = align_side

        for c_idx, day in enumerate(x_headers, start=4):
            val = row_tree[(t_block, j_name, t_slot, s_idx)].get(day, "")
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font, cell.fill, cell.border, cell.alignment = font_body, fill_white, border_cell, align_center

    end_row = start_row + len(sorted_row_keys) - 1

    if sorted_row_keys and getattr(blueprint, "auto_merge_parent_cells", True):
        merge_start = start_row
        for r in range(start_row + 1, end_row + 2):
            prev_val = ws.cell(row=r - 1, column=1).value
            curr_val = ws.cell(row=r, column=1).value if r <= end_row else None
            if curr_val != prev_val:
                if r - 1 > merge_start:
                    ws.merge_cells(start_row=merge_start, end_row=r - 1, start_column=1, end_column=1)
                ws.cell(row=merge_start, column=1).alignment = align_center
                merge_start = r

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


# --- RENDER ANALYTICS SHEET ---

def _render_analytics_sheet(wb, metrics: dict, is_rtl: bool):
    ws = wb.create_sheet(title="אנליטיקה" if is_rtl else "Analytics")
    ws.views.sheetView[0].rightToLeft = is_rtl

    try:
        ws.sheet_view.showGridLines = True
    except AttributeError:
        pass

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="1F4E78")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_body = Font(name="Calibri", size=10)

    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_card = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")

    thin_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    h_align = "right" if is_rtl else "left"
    align_center = Alignment(horizontal="center", vertical="center")
    align_side = Alignment(horizontal=h_align, vertical="center")

    current_row = 1

    summary_cards = metrics.get("summary_cards", [])
    manpower = metrics.get("manpower_utilization", [])

    # Calculate Avg Hours per person if missing from summary_cards
    if manpower and not any("avg hours" in str(card.get("label", "")).lower() or "ממוצע שעות" in str(card.get("label", "")) for card in summary_cards):
        total_hours = sum(w.get("assigned_hours", 0) for w in manpower)
        avg_hours = round(total_hours / len(manpower), 2) if manpower else 0.0
        
        card_label = "ממוצע שעות לעובד" if is_rtl else "Avg Hours / Person"
        summary_cards.append({"label": card_label, "value": f"{avg_hours} hrs"})

    if summary_cards:
        ws.cell(row=current_row, column=1, value="מדדי כוח אדם כלליים" if is_rtl else "Overall Manpower Pool Metrics").font = font_section
        current_row += 1
        for card in summary_cards:
            c1 = ws.cell(row=current_row, column=1, value=card.get("label", ""))
            c2 = ws.cell(row=current_row, column=2, value=str(card.get("value", "")))
            c1.font, c2.font = font_body, font_bold
            c1.fill, c2.fill = fill_card, fill_card
            c1.border, c2.border = border_cell, border_cell
            c1.alignment, c2.alignment = align_side, align_center
            current_row += 1
        current_row += 2

    jobs = metrics.get("job_coverage", [])
    if jobs:
        ws.cell(row=current_row, column=1, value="Job Category Performance (TCU & CDR)").font = font_section
        current_row += 1
        headers = ["Job Category", "Required Slots", "Filled Slots", "Job Utilization (TCU)", "Avg. Job Despair (CDR)"]
        for col_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=h_text)
            c.font, c.fill, c.alignment = font_header, fill_header, align_center
        current_row += 1
        for j in jobs:
            vals = [
                j.get("job_name", j.get("job_id", "")),
                j.get("required_slots", 0),
                j.get("filled_slots", 0),
                str(j.get("job_tcu", "0.0%")),
                str(j.get("job_avg_cdr", "0.00"))
            ]
            for col_idx, val in enumerate(vals, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font, c.border = font_body, border_cell
                c.alignment = align_side if col_idx == 1 else align_center
            current_row += 1
        current_row += 2

    if manpower:
        ws.cell(row=current_row, column=1, value="Individual Worker Capacity & Despair").font = font_section
        current_row += 1
        headers = ["Worker ID", "Name", "Assigned Hrs", "Max Capacity Hrs", "Worker TCU", "Worker CDR"]
        for col_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=h_text)
            c.font, c.fill, c.alignment = font_header, fill_header, align_center
        current_row += 1
        for w in manpower:
            vals = [
                w.get("worker_id", ""),
                w.get("name", ""),
                w.get("assigned_hours", 0),
                w.get("max_capacity_hours", 0),
                str(w.get("tcu_percent", "0.0%")),
                str(w.get("cdr_score", "0.00"))
            ]
            for col_idx, val in enumerate(vals, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=val)
                c.font, c.border = font_body, border_cell
                c.alignment = align_side if col_idx in (1, 2) else align_center
            current_row += 1

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 18)


# --- RENDER DIAGNOSTICS SHEET ---

def _render_diagnostics_sheet(wb, diagnostics: list, is_rtl: bool):
    ws = wb.create_sheet(title="דיאגנוסטיקה" if is_rtl else "System_Diagnostics")
    ws.views.sheetView[0].rightToLeft = is_rtl

    try:
        ws.sheet_view.showGridLines = True
    except AttributeError:
        pass

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10)
    fill_header = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    
    thin_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    h_align = "right" if is_rtl else "left"
    align_center = Alignment(horizontal="center", vertical="center")
    align_side = Alignment(horizontal=h_align, vertical="center", wrap_text=True)

    headers = ["Severity", "Code", "Affected Entity / Role", "Diagnostic Reason", "Actionable Remediation"]
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.font, cell.fill, cell.alignment = font_header, fill_header, align_center

    for r_idx, d in enumerate(diagnostics, start=2):
        c1 = ws.cell(row=r_idx, column=1, value=d.get("severity", "INFO"))
        c2 = ws.cell(row=r_idx, column=2, value=d.get("code", "LOG"))
        c3 = ws.cell(row=r_idx, column=3, value=d.get("entity", "System"))
        c4 = ws.cell(row=r_idx, column=4, value=d.get("reason", d.get("description", "")))
        c5 = ws.cell(row=r_idx, column=5, value=d.get("remediation", ""))

        for cell in (c1, c2, c3, c4, c5):
            cell.font, cell.border = font_body, border_cell

        c1.alignment, c2.alignment = align_center, align_center
        c3.alignment, c4.alignment, c5.alignment = align_side, align_side, align_side

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 45
    ws.column_dimensions['E'].width = 45


# --- RENDER TABULAR LAYOUT ---

def _render_tabular_layout(ws, assignments: List[CanonicalAssignment], blueprint: LayoutBlueprint, is_rtl: bool):
    headers = ["Shift ID", "Day", "Block Label", "Time Slot", "Job", "Worker UID", "Worker Name"]
    ws.append(headers)
    for a in assignments:
        ws.append([
            a.shift_id,
            a.day,
            a.time_block,
            a.time_slot,
            a.job_name,
            a.worker_uid,
            a.worker_name
        ])